#!/usr/bin/env python3
"""Ubah keyword plan (.xlsx hasil skill keyword-research) jadi data/brief.json.

Pemakaian:  python3 tools/brief_to_json.py <keyword-plan.xlsx> data/brief.json
Butuh openpyxl (pip install openpyxl).

Yang diambil:
- Sheet "Page Brief"  -> satu entri per halaman (title, meta, H1, keyword, link wajib)
- Sheet "Keyword Plan" -> daftar keyword per slug (untuk cek kanibalisasi)
- Sheet "Negative Keywords (Ads)" kategori "Nama kompetitor" -> daftar kata terlarang di halaman
"""
import json
import re
import sys

import openpyxl


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        if any(c is not None for c in r):
            yield head, r


def split_list(v):
    if not v:
        return []
    return [s.strip() for s in re.split(r"[;\n]", str(v)) if s.strip() and s.strip() != "—"]


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    pages = []
    ws = wb["Page Brief"]
    for head, r in rows(ws):
        # Kolom Char & Cek muncul dua kali (Title & Meta) -> baca berdasarkan posisi nama unik
        d = {}
        for h, v in zip(head, r):
            d.setdefault(h, v)
        slug = d.get("URL Slug")
        if not slug or not str(slug).startswith("/"):
            continue  # baris TOTAL
        pages.append({
            "slug": str(slug).strip(),
            "priority": d.get("Prioritas"),
            "name": d.get("Nama Halaman"),
            "ad_group": d.get("Ad Group"),
            "title": str(d.get("Title Tag")).strip(),
            "meta": str(d.get("Meta Description")).strip(),
            "h1": str(d.get("H1")).strip(),
            "primary": str(d.get("Primary Keyword")).strip(),
            "secondary": [s for s in split_list(d.get("Secondary Keyword")) if not s.startswith("(")],
            "must_have": [s.strip() for s in str(d.get("Isi Wajib di Halaman") or "").split("·") if s.strip()],
            "links": split_list(d.get("Internal Link ke")),
        })

    keywords = {}
    for head, r in rows(wb["Keyword Plan"]):
        d = dict(zip(head, r))
        if d.get("URL Slug") and d.get("Keyword"):
            keywords.setdefault(d["URL Slug"], []).append(str(d["Keyword"]))
    for p in pages:
        p["keywords"] = keywords.get(p["slug"], [])

    competitors = []
    for name in wb.sheetnames:
        if name.lower().startswith("negative"):
            for head, r in rows(wb[name]):
                d = dict(zip(head, r))
                if str(d.get("Kategori", "")).lower().startswith("nama kompetitor"):
                    competitors.append(str(d["Negative Keyword"]).strip())

    out = {"source": src.split("/")[-1], "pages": pages, "competitors": competitors}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"{len(pages)} halaman, {sum(len(v) for v in keywords.values())} keyword, "
          f"{len(competitors)} nama kompetitor -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
